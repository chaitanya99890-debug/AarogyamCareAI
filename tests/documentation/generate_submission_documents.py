"""Build the completed AarogyamCareAI testing portfolio as DOCX and PDF."""

from __future__ import annotations

from pathlib import Path
from html import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "AarogyamCareAI-Testing-Documentation-Portfolio"
SOURCES = [
    ("Software Test Plan", BASE / "software-test-plan.md"),
    ("Manual Functional Test Suite", BASE / "manual-functional-test-suite.md"),
    ("Selenium Test Execution Report", BASE / "selenium-test-execution-report.md"),
    ("Defect Report", BASE / "defect-report.md"),
    ("Test Summary Report", BASE / "test-summary-report.md"),
]


def blocks(path: Path):
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            code = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            yield ("code", "\n".join(code))
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = [
                [cell.strip() for cell in row.strip().strip("|").split("|")]
                for row in table_lines
                if not set(row.replace("|", "").strip()) <= {"-", ":"}
            ]
            if rows:
                yield ("table", rows)
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            yield (f"heading{level}", line[level:].strip())
        elif line.startswith("- "):
            yield ("bullet", line[2:].strip())
        else:
            yield ("paragraph", line.replace("  ", " ").strip())
        index += 1


def clean(text: str) -> str:
    return text.replace("**", "").replace("`", "").replace("[", "").replace("]", "")


def add_docx_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(row):
            cells[column].text = clean(value)
            if row_index == 0:
                for run in cells[column].paragraphs[0].runs:
                    run.bold = True


def build_docx() -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    title = document.add_heading("AarogyamCareAI", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Software Testing Documentation Portfolio")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Prepared for final-year project submission | 20 October 2026").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    for source_index, (name, path) in enumerate(SOURCES):
        if source_index:
            document.add_page_break()
        document.add_heading(name, 0)
        for kind, content in blocks(path):
            if kind.startswith("heading"):
                document.add_heading(clean(content), min(int(kind[-1]), 4))
            elif kind == "paragraph":
                document.add_paragraph(clean(content))
            elif kind == "bullet":
                document.add_paragraph(clean(content), style="List Bullet")
            elif kind == "code":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(content)
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            elif kind == "table":
                add_docx_table(document, content)
    output = OUTPUT.with_suffix(".docx")
    document.save(output)
    return output


def build_pdf() -> Path:
    output = OUTPUT.with_suffix(".pdf")
    document = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=1.25 * cm, leftMargin=1.25 * cm, topMargin=1.25 * cm, bottomMargin=1.25 * cm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PortfolioTitle", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=16))
    styles.add(ParagraphStyle(name="CodeBlock", parent=styles["Code"], fontSize=7, leading=9, backColor=colors.whitesmoke, borderPadding=5, spaceAfter=7))
    story = [Paragraph("AarogyamCareAI", styles["PortfolioTitle"]), Paragraph("Software Testing Documentation Portfolio", styles["Title"]), Spacer(1, 12), Paragraph("Prepared for final-year project submission | 20 October 2026", styles["Normal"]), PageBreak()]
    for source_index, (name, path) in enumerate(SOURCES):
        if source_index:
            story.append(PageBreak())
        story.append(Paragraph(escape(name), styles["Title"]))
        for kind, content in blocks(path):
            if kind.startswith("heading"):
                story.append(Paragraph(escape(clean(content)), styles[f"Heading{min(int(kind[-1]), 3)}"]))
            elif kind == "paragraph":
                story.append(Paragraph(escape(clean(content)), styles["BodyText"]))
            elif kind == "bullet":
                story.append(Paragraph("• " + escape(clean(content)), styles["BodyText"]))
            elif kind == "code":
                story.append(Paragraph(escape(content).replace("\n", "<br/>"), styles["CodeBlock"]))
            elif kind == "table":
                rows = [[Paragraph(escape(clean(value)), styles["BodyText"]) for value in row] for row in content]
                table = Table(rows, repeatRows=1)
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEBF7")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]))
                story.append(table)
                story.append(Spacer(1, 6))
    document.build(story)
    return output


if __name__ == "__main__":
    print(build_docx())
    print(build_pdf())
